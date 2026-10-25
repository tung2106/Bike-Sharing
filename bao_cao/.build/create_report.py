from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'Bao_cao_Bike_Sharing.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.75)
s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'Báo cáo tổng kết dự án Bike Sharing'
doc.core_properties.subject = 'Tiến độ thực hiện và giải thích quy trình dữ liệu mô hình website'
doc.core_properties.author = 'Dự án Bike Sharing'
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('Bike Sharing  |  ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate([headers]+rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DFEAF2' if i == 0 else ('F5F7F9' if i%2 == 0 else 'FFFFFF')); tcPr.append(shade)
            mar = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                elem = OxmlElement('w:'+side); elem.set(qn('w:w'),'95'); elem.set(qn('w:type'),'dxa'); mar.append(elem)
            tcPr.append(mar)
            pp = cell.paragraphs[0]; pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.05
            rr = pp.add_run(str(text)); rr.font.size = Pt(10); rr.bold = i == 0
        pr = t.rows[i]._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0: pr.append(OxmlElement('w:tblHeader'))
    p('').paragraph_format.space_after = Pt(0)
    return t

p('Báo cáo tổng kết\ndự án Bike Sharing', 'Title')
p('Kết quả thực hiện và giải thích quy trình', 'Subtitle')
p('Ngày tổng hợp 10 tháng 10 năm 2026')
h('Kết luận chính')
p('Dự án đã hoàn thành một quy trình học máy dự đoán số lượt thuê xe: chuẩn bị dữ liệu, chia tập theo thời gian, tiền xử lý, huấn luyện Random Forest, chọn tham số, đánh giá và đưa mô hình theo giờ vào website chạy trên máy tính. Dữ liệu theo ngày và theo giờ đều đã có kết quả huấn luyện trong notebook Colab đã lưu.')
p('Mô hình theo giờ đạt RMSE 90,16 lượt, MAE 61,36 lượt và R² 0,8243 trên tập test. Sáu kiểm thử website được chạy lại ngày 10/10/2026 đều đạt; các chỉ số test được tính lại từ mô hình hiện có và khớp kết quả đã lưu. Hệ thống hiện phù hợp để học tập và trình diễn với bối cảnh dữ liệu năm 2011–2012.')
table(['Hạng mục', 'Trạng thái hiện tại'], [
    ['Dữ liệu và tiền xử lý', 'Có dữ liệu gốc, 6 tập CSV đã xử lý và báo cáo chia tập.'],
    ['Huấn luyện và đánh giá', 'Đã chạy day và hour trên Colab; đã chọn trong 8 cấu hình mỗi bộ.'],
    ['Website dự đoán', 'Có Flask API, giao diện tiếng Việt và mô hình hour thật.'],
    ['Đưa lên Internet', 'Cấu hình hiện tại chỉ phục vụ tại 127.0.0.1:8000.'],
    ['Dự báo hiện tại và vận hành thực tế', 'Chưa có dữ liệu mới, thời tiết tự động hay giám sát mô hình.'],
], [2.05, 4.95])
h('Nội dung báo cáo')
p('Các phần tiếp theo giải thích công nghệ và cấu trúc dự án; ý nghĩa dữ liệu; quá trình làm sạch và chia tập; hai cách tiền xử lý; cơ chế huấn luyện; kết quả; cách lưu mô hình; luồng website; kiểm thử; phần còn thiếu và cách sử dụng lại.')
p('Cách đọc trạng thái: “đã hoàn thành” có mã nguồn hoặc đầu ra cụ thể; “đã kiểm chứng lại” có kết quả chạy trong lần rà soát này; các hạng mục mở rộng được trình bày riêng ở phần 11.')

page('1 Công nghệ và cấu trúc dự án')
p('Dự án gồm ba khối chính: xử lý dữ liệu bằng Python, huấn luyện trên Google Colab và dự đoán qua website. Website tải mô hình đã học sẵn, không huấn luyện lại khi người dùng bấm dự đoán.')
table(['Thành phần', 'Đã dùng vào việc gì'], [
    ['Python', 'Điều phối xử lý CSV, huấn luyện, lưu mô hình và máy chủ dự đoán.'],
    ['csv và thư viện chuẩn', 'Script preprocess.py làm sạch, chia tập và one-hot mà không cần pandas.'],
    ['pandas và NumPy', 'Bảng dữ liệu, chuyển kiểu số, giá trị thiếu và tính toán trong Colab/API.'],
    ['scikit-learn', 'SimpleImputer, OneHotEncoder, ColumnTransformer, Pipeline, RandomForestRegressor và chỉ số đánh giá.'],
    ['Google Colab và Drive', 'Chạy notebook bằng CPU; lưu dữ liệu và kết quả theo từng RUN_ID.'],
    ['joblib', 'Lưu và tải lại pipeline đã học cùng thông tin phiên bản.'],
    ['Flask và Waitress', 'Flask định nghĩa trang/API; Waitress phục vụ ứng dụng tại cổng 8000.'],
    ['HTML CSS JavaScript', 'Biểu mẫu, giao diện thích ứng, gửi JSON và hiển thị kết quả.'],
    ['unittest và Playwright', 'Kiểm thử API/mô hình và mã kiểm tra giao diện trình duyệt.'],
], [1.85, 5.15])
h('Các thư mục cần biết')
p('du_lieu/goc chứa day.csv, hour.csv và Readme.txt. du_lieu/da_xu_ly chứa đầu ra của script tiền xử lý ban đầu. tien_xu_ly/preprocess.py là mã tạo lại những đầu ra này.')
p('colab chứa notebook, mã huấn luyện, script tạo notebook và script kiểm tra. colab/ket_qua/Colab_shared_review.ipynb giữ các kết quả chạy đã lưu. website chứa API, giao diện, mô hình, hướng dẫn và các kiểm thử.')
p('ban_dong_goi chứa Bike_Sharing_Colab.zip và Velo_Hour_Website.zip để chuyển sang môi trường khác. MO_WEBSITE.cmd là điểm mở website bằng nhấp đúp. Các gói ZIP là bản đóng gói tại một thời điểm, cần tạo lại khi mã hoặc mô hình thay đổi.')

page('2 Dữ liệu đầu vào và bài toán')
p('Bài toán là hồi quy: từ điều kiện lịch và thời tiết, ước lượng một con số là tổng lượt thuê xe cnt. Dữ liệu gốc mô tả Capital Bikeshare tại Washington D.C. trong hai năm 2011 và 2012. day.csv có 731 dòng theo ngày; hour.csv có 17.379 dòng theo giờ. Hai bộ được xử lý và huấn luyện riêng.')
table(['Nhóm cột', 'Ý nghĩa và cách sử dụng'], [
    ['cnt', 'Nhãn cần dự đoán; tổng lượt thuê trong một ngày hoặc một giờ.'],
    ['temp và atemp', 'Nhiệt độ và nhiệt độ cảm nhận; dữ liệu CSV đã ở thang chuẩn hóa.'],
    ['hum và windspeed', 'Độ ẩm và tốc độ gió đã được chuẩn hóa trong dữ liệu gốc.'],
    ['yr holiday workingday', 'Biến nhị phân: năm 2011/2012, ngày lễ, ngày làm việc.'],
    ['season mnth weekday', 'Mã mùa, tháng và thứ trong tuần; dùng như biến phân loại.'],
    ['weathersit', 'Mã nhóm thời tiết từ 1 đến 4; giá trị lớn hơn mô tả điều kiện xấu hơn.'],
    ['hr', 'Giờ từ 0 đến 23; chỉ có trong hour.csv.'],
    ['instant và dteday', 'Mã bản ghi và ngày; phục vụ kiểm tra/chia thời gian, không đưa trực tiếp vào mô hình.'],
    ['casual và registered', 'Lượt thuê khách vãng lai và thành viên; bị loại khỏi đầu vào.'],
], [2, 5])
h('Tại sao bỏ casual và registered')
p('Trong dữ liệu, cnt là tổng casual và registered. Nếu đưa hai cột này vào mô hình, mô hình gần như nhìn thấy đáp án. Đây là rò rỉ nhãn: kết quả kiểm tra có thể rất tốt nhưng không phản ánh khả năng dự đoán trước khi biết lượt thuê thực tế.')
h('Phạm vi dữ liệu cần hiểu đúng')
p('Không có ô trống không đồng nghĩa chuỗi giờ liên tục. Hai năm gồm 731 ngày tương ứng 17.544 giờ, nhiều hơn 17.379 bản ghi hiện có 165 giờ. Quy trình hiện giữ các bản ghi quan sát được, không tự chèn giờ vắng mặt hoặc gán chúng bằng 0. Kết quả dự đoán là tổng nhu cầu, không phải số xe còn sẵn tại từng trạm.')

page('3 Làm sạch dữ liệu trước khi học')
h('Bước 1 Đọc và kiểm tra cấu trúc')
p('Chương trình kiểm tra các cột bắt buộc trước khi xử lý. Các ký hiệu thiếu như ô rỗng, NA, N/A, NaN, null, none và dấu hỏi được nhận diện. Ở luồng Colab, pandas đọc CSV, normalize_features chuyển các đặc trưng thành số và chuyển giá trị vô hạn thành giá trị thiếu.')
h('Bước 2 Kiểm tra bản ghi và nhãn')
p('Các dòng trùng hoàn toàn được loại. Dòng không có cnt bị loại thay vì tự tạo nhãn. cnt phải là số nguyên không âm; biến nhị phân phải bằng 0 hoặc 1; mã phân loại phải là số nguyên. Ngày không hợp lệ, giờ thiếu hoặc ngoài khoảng 0–23 làm chương trình báo lỗi vì không thể xác định đúng thứ tự thời gian.')
h('Bước 3 Sắp xếp và phát hiện xung đột')
p('Dữ liệu ngày được sắp theo ngày; dữ liệu giờ được ghép ngày với hr để tạo timestamp rồi sắp tăng dần. Script ban đầu kiểm tra trùng timestamp; luồng Colab còn kiểm tra trùng instant. Bản ghi xung đột cần được kiểm tra nguồn, không được im lặng lấy một dòng bất kỳ.')
h('Kết quả làm sạch dữ liệu hiện có')
table(['Chỉ tiêu trong report.json', 'day', 'hour'], [
    ['Số dòng ban đầu', '731', '17.379'],
    ['Số dòng giữ lại', '731', '17.379'],
    ['Dòng trùng bị loại', '0', '0'],
    ['Dòng thiếu cnt bị loại', '0', '0'],
    ['Ô thiếu trong dữ liệu gốc được báo cáo', '0', '0'],
    ['Ô đặc trưng cần điền trong các tập đã xuất', '0', '0'],
], [4.4, 1.3, 1.3])
p('Như vậy, dự án đã xây dựng cơ chế xử lý thiếu và trùng, nhưng dữ liệu đang dùng không cần thực sự loại dòng hay điền ô thiếu. Cần phân biệt việc có cơ chế xử lý với việc cơ chế đó đã thay đổi bao nhiêu dữ liệu.')
p('Quy trình chưa triển khai phát hiện và xử lý ngoại lệ thống kê chuyên biệt. Những lượt thuê rất cao hoặc rất thấp vẫn được giữ lại nếu thỏa điều kiện hợp lệ. Cũng chưa có quy trình bổ sung dữ liệu thời gian thực hoặc suy đoán nhãn cho các giờ vắng mặt.')

page('4 Chia tập và hai luồng tiền xử lý')
p('Dữ liệu được chia theo thời gian, không xáo trộn. Khoảng 70% đầu dùng học, 15% tiếp theo dùng chọn cấu hình, phần còn lại dùng đánh giá cuối. Số dòng train và validation được làm tròn xuống nên tỷ lệ không đúng tuyệt đối 70/15/15.')
table(['Bộ và tập', 'Số dòng', 'Khoảng thời gian'], [
    ['day train', '511', '01/01/2011 đến 25/05/2012'],
    ['day validation', '109', '26/05/2012 đến 11/09/2012'],
    ['day test', '111', '12/09/2012 đến 31/12/2012'],
    ['hour train', '12.165', '01/01/2011 00:00 đến 27/05/2012 01:00'],
    ['hour validation', '2.606', '27/05/2012 02:00 đến 12/09/2012 15:00'],
    ['hour test', '2.608', '12/09/2012 16:00 đến 31/12/2012 23:00'],
], [1.5, .8, 4.7])
h('Chỉ học cách xử lý từ train')
p('Biến số được điền bằng trung vị; biến phân loại và nhị phân bằng giá trị xuất hiện nhiều nhất. Các giá trị thay thế và danh sách danh mục chỉ được học từ train. Validation, test và đầu vào mới chỉ áp dụng lại các thông số đó. Nhờ vậy, thông tin ở phần tương lai không ảnh hưởng ngược tới bước học.')
p('One-hot biến một mã thành nhiều cột 0/1. Ví dụ, tháng 5 làm cột mnth_5 bằng 1 và các cột tháng còn lại bằng 0. Biến nhị phân giữ nguyên; thời tiết trong CSV đã chuẩn hóa nên không được scale lại trong bước học.')
table(['Điểm khác nhau', 'Script ban đầu', 'Pipeline Colab và website'], [
    ['Cách thực hiện', 'Thư viện chuẩn Python', 'scikit-learn'],
    ['Danh mục chưa gặp', 'Có cột _unknown riêng', 'Nhóm one-hot tương ứng toàn 0'],
    ['Số đặc trưng sau mã hóa', 'day 37; hour 63', 'day 33; hour 58'],
    ['Đầu ra sử dụng', 'CSV và preprocessor.json', 'Pipeline đã học trong model.joblib'],
], [2, 2.5, 2.5])
p('Số đặc trưng không tính cột cnt. Website nhận các đặc trưng gốc và tự chạy pipeline, không nhận CSV đã one-hot. Không thay thế dữ liệu của hai luồng cho nhau chỉ vì chúng cùng có tên train.csv.')

page('5 Huấn luyện và chọn mô hình')
h('Random Forest hoạt động như thế nào')
p('Random Forest là tập hợp nhiều cây quyết định. Mỗi cây học những quy tắc chia dữ liệu, chẳng hạn tách theo giờ, thời tiết và điều kiện lịch để tạo các nhóm có nhu cầu thuê gần nhau. Với thiết lập đang dùng, mỗi cây được học từ một mẫu lấy lại có hoàn lại của train. Kết quả cuối là trung bình dự đoán của các cây, nên có thể là số thập phân.')
p('ColumnTransformer áp dụng cách xử lý phù hợp cho từng nhóm cột. SimpleImputer điền thiếu; OneHotEncoder mã hóa phân loại; RandomForestRegressor học quan hệ giữa đặc trưng và cnt. Pipeline ghép bộ xử lý và mô hình đã học thành một đối tượng để dùng lại thống nhất.')
h('Tám cấu hình được so sánh cho mỗi bộ dữ liệu')
table(['Tham số', 'Giá trị thử', 'Tác dụng'], [
    ['n_estimators', '100 hoặc 200', 'Số cây trong rừng.'],
    ['max_depth', 'None hoặc 16', 'Giới hạn độ sâu; None không đặt giới hạn này.'],
    ['min_samples_leaf', '1 hoặc 3', 'Số mẫu tối thiểu ở một nút lá.'],
    ['max_features', '1.0', 'Xét toàn bộ đặc trưng khi tìm cách chia nút.'],
    ['random_state', '42 cố định', 'Giúp tái lập việc lấy mẫu ngẫu nhiên.'],
], [2, 1.5, 3.5])
p('ParameterGrid tạo 2 × 2 × 2 × 1 = 8 tổ hợp. Bộ tiền xử lý được fit một lần trên train; từng Random Forest học trên train đã mã hóa và được đo RMSE trên validation. Cấu hình có RMSE validation nhỏ nhất được giữ lại; nếu bằng nhau, mã hiện tại giữ cấu hình gặp trước.')
h('Cấu hình cuối đã chọn')
p('day: 200 cây, độ sâu tối đa 16, min_samples_leaf = 1, max_features = 1.0. RMSE validation đạt 834,4331. hour: 100 cây, không giới hạn max_depth, min_samples_leaf = 1, max_features = 1.0. RMSE validation đạt 126,3800.')
p('Sau khi chọn, dự án không huấn luyện lại trên train cộng validation. Mô hình cuối vẫn chỉ học từ 70% train. Test chỉ được đưa vào bước đánh giá sau khi đã chốt cấu hình. Đây là một lần chia theo thời gian, chưa phải kiểm định qua nhiều cửa sổ thời gian.')

page('6 Kết quả thực tế và cách diễn giải')
table(['Mô hình', 'RMSE test', 'MAE test', 'R² test', 'RMSE mốc'], [
    ['day', '1.362,3345', '1.209,6837', '0,5376', '2.679,2468'],
    ['hour', '90,1641', '61,3564', '0,8243', '228,1061'],
], [1, 1.6, 1.6, 1.2, 1.6])
p('Nguồn kết quả: bảng đầu ra trong Colab_shared_review.ipynb cho cả day và hour; metrics.json của website lưu chi tiết hour. Mốc so sánh là cách luôn dự đoán bằng trung bình cnt của train, không sử dụng thông tin thời tiết hoặc lịch của từng mẫu.')
h('MAE cho biết sai lệch trung bình')
p('MAE lấy trị tuyệt đối của từng sai số rồi tính trung bình. Với hour, MAE 61,36 nghĩa là trên tập test, dự đoán lệch trung bình khoảng 61 lượt thuê mỗi bản ghi giờ. Đây không phải cam kết rằng mọi lần dự đoán đều nằm trong khoảng cộng trừ 61 lượt.')
h('RMSE nhấn mạnh những sai số lớn')
p('RMSE bình phương sai số, tính trung bình rồi lấy căn bậc hai. Vì sai số lớn bị phạt mạnh hơn, RMSE thường cao hơn MAE. Mô hình hour giảm RMSE khoảng 60,47% so với mốc trung bình train; day giảm khoảng 49,15%. Các tỷ lệ này được tính từ kết quả test đã lưu.')
h('R² không phải phần trăm dự đoán đúng')
p('R² so sánh tổng bình phương sai số với mức biến thiên của nhãn trên tập được đánh giá. R² của hour bằng 0,8243 cho thấy mô hình giải thích được khoảng 82,43% biến thiên cnt trong tập test theo cách đo này. Không nên diễn đạt là “độ chính xác 82,43%”. R² có thể âm nếu dự đoán kém hơn mốc trung bình nhãn của chính tập đánh giá.')
h('Những giới hạn khi đọc kết quả')
p('Không so RMSE day trực tiếp với RMSE hour để kết luận bộ nào tốt hơn: đơn vị tổng hợp là lượt/ngày và lượt/giờ nên độ lớn nhãn khác nhau. Mỗi mô hình cần được so với mốc cùng tập và cùng đơn vị.')
p('day có R² validation khoảng 0,0385, trong khi R² test khoảng 0,5376. hour có R² validation khoảng 0,6795 và test khoảng 0,8243. Hiệu quả thay đổi giữa các giai đoạn là lý do cần đánh giá qua nhiều cửa sổ thời gian. Kết quả hiện có chưa chứng minh khả năng áp dụng cho năm 2026 hay cho một thành phố khác.')

page('7 Lưu kết quả và sử dụng lại mô hình')
h('Các đầu ra của một lần chạy Colab')
p('Kết quả được tổ chức dưới MyDrive/TT_CSN/runs/<RUN_ID>/, tách day và hour. Notebook đã lưu thể hiện lần chạy 20261004_014306_709f7fb3. Mỗi lần thử nghiệm cần thư mục mới để tách kết quả, cấu hình và dữ liệu của các lần chạy.')
table(['Tệp hoặc nhóm tệp', 'Nội dung và mục đích'], [
    ['splits/raw/*.csv', 'Đặc trưng gốc và cnt sau chia tập, trước điền thiếu/one-hot.'],
    ['splits/*.csv', 'Ma trận đặc trưng đã mã hóa và cnt.'],
    ['split_report.json và split_manifest.csv', 'Kích thước tập, khoảng thời gian và nguồn gốc bản ghi; báo cáo Colab có mã băm dữ liệu nguồn.'],
    ['validation_results.csv và selection.json', 'Điểm của các cấu hình, cấu hình tốt nhất và tiêu chí lựa chọn.'],
    ['metrics.json và test_predictions.csv', 'Điểm test, mốc so sánh, giá trị thật/dự đoán và residual = thực tế trừ dự đoán.'],
    ['model.joblib', 'Bundle gồm pipeline, features, target, dataset, versions và report.'],
    ['preprocessor.joblib', 'Bộ tiền xử lý đã fit; xuất riêng để kiểm tra hoặc dùng lại.'],
    ['environment.json và requirements.txt', 'Phiên bản môi trường và thư viện của lần huấn luyện.'],
], [2.9, 4.1])
h('Vì sao phải lưu cả pipeline')
p('Nếu chỉ lưu Random Forest nhưng tự mã hóa lại đầu vào, thứ tự cột hoặc giá trị điền thiếu có thể khác lúc học. Lưu cả pipeline bảo đảm đầu vào mới đi qua đúng các phép biến đổi đã học. Hàm predict_saved_model nhận DataFrame hoặc CSV có các cột gốc, không cần cnt.')
h('Mô hình website đang dùng')
p('website/models/model.joblib là bundle hour đã tải về, không huấn luyện lại trong lúc mở web. environment.json ghi Python 3.13.15, NumPy 2.1.3, pandas 2.2.3, scikit-learn 1.6.1 và joblib 1.6.0 của môi trường huấn luyện. Website kiểm tra loại dataset, thứ tự cột, nhãn và phiên bản scikit-learn khi tải.')
p('Mã băm SHA-256 của model.joblib đã được đối chiếu và khớp provenance.json. Mã băm giúp nhận diện chính xác tệp đang sử dụng; nó không phải chỉ số chất lượng dự đoán. Các kết quả day hiện được thể hiện trong notebook đã lưu, còn thư mục models của website chỉ chứa bundle hour.')

page('8 Website hoạt động từ đầu vào đến kết quả')
h('Luồng xử lý một yêu cầu')
p('Khi mở trang, JavaScript gọi GET /api/model để biết mô hình đã sẵn sàng chưa và lấy các chỉ số test. Người dùng chọn ngày, giờ, nhóm thời tiết, trạng thái ngày lễ và nhập các thông số thời tiết. Khi bấm tính, giao diện gửi JSON đến POST /api/predict.')
p('Flask kiểm tra đầu vào, suy ra các biến lịch, chuyển đơn vị thời tiết về thang dữ liệu và tạo một DataFrame gồm 12 đặc trưng theo đúng thứ tự. Pipeline điền thiếu và one-hot, sau đó Random Forest trả dự đoán. API gửi kết quả dạng số thực và số làm tròn; giao diện hiển thị lượt thuê ước tính.')
h('Các biến lịch được tự tạo')
p('yr là năm trừ 2011; mnth là tháng; weekday dùng Chủ nhật = 0. workingday bằng 1 khi là thứ Hai đến thứ Sáu và không được đánh dấu ngày lễ. Người dùng tự chọn ngày lễ; chương trình chưa tự tra lịch nghỉ lễ. API chỉ chấp nhận ngày thuộc 2011 hoặc 2012.')
p('season dùng các mốc trong mã hiện tại: 21/03, 21/06, 23/09 và 21/12. Kiểm thử đối chiếu những mã này với tất cả ngày trong hour.csv. Cách gọi tên mùa trên giao diện và mô tả nhãn mùa trong Readme gốc cần được thống nhất khi hoàn thiện tài liệu.')
h('Ví dụ đã chạy qua API với mô hình thật')
table(['Đầu vào người dùng', 'Giá trị'], [
    ['Ngày và giờ', '01/05/2012 lúc 08:00'],
    ['Thời tiết và ngày lễ', 'Mã 1; không phải ngày lễ'],
    ['Nhiệt độ và cảm nhận', '24 °C và 26 °C'],
    ['Độ ẩm và tốc độ gió', '65% và 12 theo đơn vị giao diện'],
    ['Các biến lịch suy ra', 'yr 1; mnth 5; weekday 2; workingday 1; season 2'],
    ['Kết quả API', '459,67 lượt; giao diện làm tròn thành 460 lượt'],
], [3.15, 3.85])
p('Ví dụ này minh họa luồng hoạt động, không phải một phép đánh giá chất lượng mới. Nhiệt độ, cảm nhận, độ ẩm và gió có thể để trống; API chuyển thành NaN để pipeline điền bằng thống kê train, đồng thời báo cho giao diện những trường đã được điền.')

page('9 Chuyển đổi đầu vào và kiểm soát lỗi')
h('Chuyển thời tiết về thang mô hình')
p('Trong website/app.py, các công thức đang dùng cho bộ hour là: temp = (nhiệt độ + 8) / 47; atemp = (nhiệt độ cảm nhận + 16) / 66; hum = độ ẩm / 100; windspeed = tốc độ gió / 67. Ví dụ ở phần trước tạo temp khoảng 0,68085, atemp khoảng 0,63636, hum 0,65 và windspeed khoảng 0,17910.')
p('Các công thức nhiệt độ này khác mô tả chia 41 và chia 50 trong Readme.txt gốc. Khi mô tả cách website đang hoạt động, cần bám theo mã app.py. Kiểm thử hiện có chuyển dữ liệu nguồn ra đơn vị giao diện rồi chuyển ngược để đối chiếu; phép kiểm tra đó xác nhận tính nhất quán của luồng, không thay cho việc xác minh độc lập tài liệu định nghĩa đơn vị.')
h('Kiểm tra phía máy chủ')
p('API yêu cầu JSON là một đối tượng; ngày hợp lệ và thuộc 2011–2012; giờ nguyên 0–23; mã thời tiết nguyên 1–4; holiday là true hoặc false. Phạm vi thời tiết được chấp nhận: nhiệt độ −8 đến 39, cảm nhận −16 đến 50, độ ẩm 0–100 và gió 0–67. Giá trị vô hạn, sai kiểu hoặc ngoài phạm vi bị từ chối.')
p('Mã HTTP 400 báo đầu vào sai; 503 báo chưa tải được mô hình; 500 báo lỗi dự đoán; 413 báo dữ liệu gửi quá lớn. Giới hạn request là 16 KB. Kết quả phải hữu hạn và không âm trước khi được trả về.')
h('Kiểm soát trạng thái phía giao diện')
p('Nút dự đoán bị khóa khi mô hình chưa sẵn sàng hoặc đang xử lý. Yêu cầu quá 30 giây được hủy phía trình duyệt. Nếu người dùng đổi đầu vào trong khi chờ, JavaScript không gắn kết quả cũ cho bộ đầu vào mới mà chuyển trạng thái sang “Cần tính lại”.')
p('Giao diện sử dụng định dạng số tiếng Việt và hiển thị các trường được điền thiếu. HTML/CSS có bố cục cho màn hình nhỏ; script Playwright được viết để kiểm tra dự đoán, lỗi nhập liệu, thay đổi đầu vào, màn hình di động và phóng to chữ.')
h('Khởi động và phạm vi phục vụ')
p('MO_WEBSITE.cmd gọi quy trình mở website; open_website.py kiểm tra trạng thái, khởi động máy chủ nếu cần rồi mở trình duyệt. Waitress phục vụ trên 127.0.0.1:8000 với 4 luồng. Địa chỉ loopback này chỉ phục vụ máy đang chạy, chưa phải địa chỉ công khai để người khác truy cập qua Internet.')

page('10 Những phần đã được kiểm chứng')
h('Kiểm thử chạy lại ngày 10 tháng 10 năm 2026')
p('Lệnh website/.venv/Scripts/python.exe -B website/test_app.py đã chạy thành công 6 kiểm thử. Đây là kiểm tra với mô hình thật trong website/models, không dùng mô hình giả và không huấn luyện lại.')
table(['Kiểm thử', 'Nội dung xác nhận'], [
    ['Trang và trạng thái', 'Trang chủ và tài nguyên trả 200; API báo mô hình sẵn sàng và có metrics.'],
    ['Dự đoán khớp pipeline', 'Kết quả API bằng kết quả gọi trực tiếp pipeline với cùng đầu vào.'],
    ['Đầu vào thiếu thời tiết', 'Bốn trường trống được điền bằng bộ xử lý đã học và cho kết quả hữu hạn.'],
    ['Đầu vào không hợp lệ', 'Từ chối ngày, giờ, kiểu boolean, thông số thời tiết và JSON không hợp lệ.'],
    ['Lịch và mùa', 'Mã thứ, ngày làm việc, ngày lễ và các mốc mùa được xử lý theo quy tắc.'],
    ['CSV và chỉ số gốc', 'Đối chiếu lịch trên dữ liệu; thử chuyển đổi một số dòng; tính lại metrics trên 2.608 dòng test.'],
], [2.05, 4.95])
p('Kết quả tính lại của hour: RMSE 90,1641035755342; MAE 61,35637480828221; R² 0,8243372986215161. Ba chỉ số khớp báo cáo trong bundle tới độ chính xác mà bộ kiểm thử yêu cầu.')
h('Bằng chứng huấn luyện trên Colab')
p('Notebook Colab_shared_review.ipynb có đầu ra kết nối Drive, số dòng các tập, kết quả của 8 cấu hình cho từng dataset, bảng test và ví dụ tải lại mô hình. Vì vậy, phần huấn luyện đã có kết quả thực tế được lưu, vượt qua trạng thái chỉ viết mã hoặc chỉ chuẩn bị notebook.')
h('Phân biệt mã kiểm thử với kết quả đã chạy lại')
p('colab/verify_colab.py có cơ chế kiểm tra notebook với lưới nhỏ và mô phỏng kết nối Drive. website/check_browser.cjs chứa kiểm thử giao diện Playwright. Hai script này không được chạy lại trong lần tổng hợp báo cáo; không xem việc có script là bằng chứng chúng vừa đạt ở môi trường hiện tại.')
p('provenance.json ghi nhận lần kiểm tra giao diện trước đây đã đạt. Trong lần tổng hợp này, việc tái kiểm chứng tập trung vào 6 kiểm thử website, đối chiếu metrics và mã băm mô hình; không thực hiện lại toàn bộ phiên huấn luyện Colab hay triển khai Internet.')

page('11 Phần chưa hoàn thiện và việc nên làm tiếp')
p('Các hạng mục dưới đây mô tả giới hạn của phiên bản hiện tại và hướng phát triển. Những chức năng sản phẩm mới chỉ cần thực hiện nếu mục tiêu dự án được mở rộng từ trình diễn học máy sang sử dụng thực tế.')
table(['Hạng mục', 'Hiện trạng và hướng xử lý'], [
    ['Tài liệu và cách tái lập', 'Ưu tiên cao. Hướng dẫn Colab còn câu “chưa có kết quả train”, trái với notebook kết quả đã lưu. Cập nhật trạng thái, đường dẫn và cách phân biệt hai luồng tiền xử lý.'],
    ['Định nghĩa dữ liệu', 'Ưu tiên cao. Thống nhất công thức nhiệt độ, đơn vị tốc độ gió và nhãn mùa giữa README, mã, giao diện và báo cáo.'],
    ['Phân tích dữ liệu và sai số', 'Chưa có phần EDA chuyên biệt trong các notebook được rà soát. Bổ sung biểu đồ theo giờ, ngày, mùa, phân phối nhãn và nhóm có sai số lớn.'],
    ['Đánh giá theo thời gian', 'Mới dùng một lần chia. Thử nhiều cửa sổ quá khứ–tương lai; chốt tiêu chí trước khi đánh giá tập cuối để tránh tối ưu theo test.'],
    ['So sánh thuật toán', 'Mới có Random Forest và mốc trung bình train. Bổ sung mô hình đối chứng phù hợp, giữ cùng cách chia và thước đo.'],
    ['Mô hình theo ngày trên web', 'day đã có kết quả trong notebook nhưng chưa được tích hợp vào website; cần bundle day, schema và giao diện tương ứng nếu muốn hỗ trợ.'],
    ['Dữ liệu và thời tiết mới', 'Chưa có nguồn cập nhật hay API thời tiết. Website giới hạn 2011–2012; cần dữ liệu phù hợp và đánh giá lại trước khi hỗ trợ năm mới hoặc địa phương mới.'],
    ['Vận hành công khai', 'Chưa có cấu hình hosting công khai, tên miền, HTTPS, kiểm thử tải hoặc giám sát chất lượng mô hình.'],
    ['Chức năng ứng dụng mở rộng', 'Chưa có tài khoản, lịch sử dự đoán, dự đoán hàng loạt, quản lý trạm/xe, đặt xe hay thanh toán. Những phần này nằm ngoài chức năng dự đoán hiện có.'],
], [1.85, 5.15])
p('Thứ tự hợp lý: hoàn thiện tài liệu và quy trình tái lập trước; phân tích sai số và kiểm định độ ổn định tiếp theo; sau đó mới bổ sung dữ liệu mới, triển khai công khai hoặc mở rộng chức năng. Chưa nên cam kết độ chính xác cho bối cảnh ngoài dữ liệu đã đánh giá.')

page('12 Cách sử dụng lại và nguồn đối chiếu')
h('Mở website đã có')
p('Trên máy hiện tại, nhấp đúp MO_WEBSITE.cmd và giữ tiến trình máy chủ hoạt động trong lúc dùng. Khi cài trên máy mới, chạy website/setup.ps1 để tạo môi trường và cài thư viện, sau đó chạy website/start.ps1. Cần có Python phù hợp và bundle hour trong website/models.')
h('Chạy lại tiền xử lý hoặc huấn luyện')
p('Từ thư mục gốc dự án, python tien_xu_ly/preprocess.py tạo lại đầu ra trong du_lieu/da_xu_ly. Lệnh này phục vụ luồng CSV ban đầu, không thay thế model.joblib của website.')
p('Để huấn luyện trên Colab: đưa hai CSV gốc vào MyDrive/TT_CSN/data, tải Bike_Sharing_Colab.ipynb lên Colab, chạy các ô từ trên xuống và cấp quyền Drive. Chọn lưới tham số trước khi xem test. Lưu mỗi thử nghiệm vào RUN_ID mới; chỉ đưa bundle hour và các tệp thông tin tương ứng sang website sau khi hoàn tất đánh giá.')
p('Để kiểm chứng website, chạy .\\website\\.venv\\Scripts\\python.exe -B website/test_app.py từ thư mục gốc. Khi thay mô hình, kiểm tra dataset, thứ tự đặc trưng, phiên bản thư viện và việc tái lập chỉ số. Không đưa CSV đã one-hot vào pipeline nhận dữ liệu gốc.')
h('Các tệp làm căn cứ cho báo cáo')
table(['Nguồn trong dự án', 'Nội dung đối chiếu'], [
    ['du_lieu/goc/Readme.txt', 'Nguồn dữ liệu, phạm vi năm, mô tả cột và thông tin trích dẫn.'],
    ['tien_xu_ly/preprocess.py; du_lieu/da_xu_ly/{day,hour}/report.json', 'Làm sạch, số dòng, thời gian chia, số đặc trưng và đầu ra CSV.'],
    ['colab/colab_train.py', 'Pipeline, lưới tham số, chọn mô hình, đánh giá và lưu bundle.'],
    ['colab/ket_qua/Colab_shared_review.ipynb', 'Kết quả thực tế của cả day và hour đã lưu từ Colab.'],
    ['website/models/metrics.json; environment.json; provenance.json', 'Chất lượng, môi trường và nguồn gốc bundle đang dùng.'],
    ['website/app.py; static/app.js; test_app.py', 'Luồng dự đoán, chuyển đổi, giao diện và kiểm thử tái lập.'],
    ['website/open_website.py; setup.ps1; start.ps1', 'Cài đặt, kiểm tra trạng thái và khởi động website.'],
], [3.6, 3.4])
p('Thông tin trích dẫn dữ liệu theo Readme đi kèm: Hadi Fanaee-T và João Gama, Event labeling combining ensemble detectors and background knowledge, Progress in Artificial Intelligence, 2013, DOI 10.1007/s13748-013-0040-3. Đây là thông tin thư mục được cung cấp cùng bộ dữ liệu.')

for node in list(doc.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
for node in list(doc.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
